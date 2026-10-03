from flask import Blueprint, request, jsonify, send_file
from sqlalchemy import func, or_
from datetime import datetime, date
from decimal import Decimal
from io import BytesIO
from collections import OrderedDict
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import db, Customer, Inventory, Movement, Order


export_bp = Blueprint("export", __name__)


# ============================================================
# CONSTANTS
# ============================================================

ORDER_STATUSES = [
    "Chờ xác nhận",
    "Đã xác nhận",
    "Đang chuẩn bị",
    "Đang làm",
    "Đã hoàn thành",
    "Đang giao",
    "Đã giao",
    "Hủy",
]

MOVEMENT_TYPES = {
    "in": "Nhập kho",
    "out": "Xuất kho",
    "waste": "Hao hụt / hỏng",
}


# ============================================================
# COMMON HELPERS
# ============================================================

def safe_value(value):
    """
    Convert SQLAlchemy / Decimal / date values
    into JSON / Excel friendly values.
    """
    if value is None:
        return ""

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return value


def format_date(value):
    """
    Excel-friendly date value.
    """
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    return value


def format_number(value):
    """
    Keep numeric values as numbers in Excel.
    """
    if value is None:
        return 0

    if isinstance(value, Decimal):
        return float(value)

    return value


def style_header(ws, row=1):
    """
    Style the first/header row.
    """
    fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    font = Font(
        bold=True,
        color="FFFFFF",
    )

    alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    thin = Side(
        style="thin",
        color="D9E1F2",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        cell.border = border

    ws.row_dimensions[row].height = 30


def style_data(ws):
    """
    Basic formatting for worksheet data.
    """
    thin = Side(
        style="thin",
        color="E5E7EB",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


def auto_width(ws, max_width=45):
    """
    Automatically calculate column widths.
    """
    for column_cells in ws.columns:
        max_length = 0

        for cell in column_cells:
            try:
                value = cell.value

                if value is None:
                    continue

                length = len(str(value))

                if length > max_length:
                    max_length = length

            except Exception:
                pass

        adjusted_width = min(
            max(max_length + 2, 10),
            max_width,
        )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        ws.column_dimensions[
            column_letter
        ].width = adjusted_width


def freeze_header(ws):
    """
    Freeze first row.
    """
    ws.freeze_panes = "A2"


def prepare_sheet(ws, headers):
    """
    Initialize a worksheet.
    """
    ws.append(headers)

    style_header(ws)

    freeze_header(ws)


def finalize_sheet(ws):
    """
    Final worksheet formatting.
    """
    style_data(ws)
    auto_width(ws)


def workbook_to_file(workbook, filename):
    """
    Convert openpyxl workbook into downloadable file.
    """
    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


def json_file_response(data, filename):
    """
    Create downloadable JSON file.
    """
    output = BytesIO()

    content = json.dumps(
        data,
        ensure_ascii=False,
        indent=2,
    )

    output.write(
        content.encode("utf-8")
    )

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/json",
    )


# ============================================================
# ORDER QUERY / FILTER
# ============================================================
def get_filtered_orders():
    """
    Get orders according to optional export filters.

    Supported query parameters:

        q
        status
        date_from
        date_to

    date_from / date_to filter theo NGÀY GIAO
    vì giao diện Orders đang sử dụng:
        "Từ ngày nhận"
        "Đến ngày nhận"
    """

    query = (
        db.session.query(Order)
        .outerjoin(
            Customer,
            Order.customer_id == Customer.id,
        )
    )

    # ========================================================
    # Keyword
    # ========================================================

    keyword = (
        request.args.get("q", "")
        .strip()
        .lower()
    )

    if keyword:
        pattern = f"%{keyword}%"

        query = query.filter(
            or_(
                func.lower(Order.id).like(pattern),
                func.lower(Order.recipient).like(pattern),
                func.lower(Order.recipient_phone).like(pattern),
                func.lower(Customer.id).like(pattern),
                func.lower(Customer.name).like(pattern),
                func.lower(Customer.phone).like(pattern),
            )
        )

    # ========================================================
    # Status
    # ========================================================

    status = (
        request.args.get("status", "")
        .strip()
    )

    if status:
        query = query.filter(
            Order.status == status
        )

    # ========================================================
    # Date from
    # ========================================================

    date_from = (
        request.args.get("date_from", "")
        .strip()
    )

    if date_from:
        try:
            parsed_date = datetime.strptime(
                date_from,
                "%Y-%m-%d",
            ).date()

            query = query.filter(
                Order.delivery_date >= parsed_date
            )

        except ValueError:
            pass

    # ========================================================
    # Date to
    # ========================================================

    date_to = (
        request.args.get("date_to", "")
        .strip()
    )

    if date_to:
        try:
            parsed_date = datetime.strptime(
                date_to,
                "%Y-%m-%d",
            ).date()

            query = query.filter(
                Order.delivery_date <= parsed_date
            )

        except ValueError:
            pass

    # ========================================================
    # Newest first
    # ========================================================

    query = query.order_by(
        Order.created.desc(),
        Order.id.desc(),
    )

    return query.all()

# ============================================================
# EXCEL - ORDERS
# ============================================================

def build_orders_sheet(
    workbook,
    orders=None,
    title="Đơn hàng",
):
    """
    Create order worksheet.
    """

    ws = workbook.create_sheet(title)

    headers = [
        "Mã đơn hàng",
        "Ngày tạo",

        "Mã khách hàng",
        "Tên khách hàng",
        "SĐT khách hàng",

        "Người nhận",
        "SĐT người nhận",

        "Ngày giao",
        "Giờ giao",
        "Địa chỉ",

        "Sản phẩm",

        "Tạm tính",
        "Phí giao hàng",

        "Loại giảm giá",
        "% giảm giá",
        "Giảm giá",
        "Giá trị giảm",

        "Tổng tiền",
        "Đã cọc",
        "Còn phải thu",

        "Thanh toán",
        "Trạng thái",

        "Thiệp",
        "Ghi chú",
    ]

    prepare_sheet(
        ws,
        headers,
    )

    if orders is None:
        orders = get_filtered_orders()

    for order in orders:

        customer = order.customer

        ws.append([
            order.id,
            format_date(order.created),

            customer.id
            if customer
            else order.customer_id,

            customer.name
            if customer
            else "",

            customer.phone
            if customer
            else "",

            order.recipient or "",
            order.recipient_phone or "",

            format_date(
                order.delivery_date
            ),

            order.delivery_time or "",
            order.address or "",

            order.product or "",

            format_number(
                order.subtotal
            ),

            format_number(
                order.shipping
            ),

            order.discount_type or "",

            format_number(
                order.discount_percent
            ),

            format_number(
                order.discount_amount
            ),

            format_number(
                order.discount_value
            ),

            format_number(
                order.total
            ),

            format_number(
                order.deposit
            ),

            format_number(
                order.balance
            ),

            order.payment or "",
            order.status or "",

            order.card or "",
            order.notes or "",
        ])

    # Date columns
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).number_format = "dd/mm/yyyy"
        ws.cell(row, 8).number_format = "dd/mm/yyyy"

    # Money columns
    money_columns = [
        12, 13, 16, 17, 18, 19, 20
    ]

    for row in range(2, ws.max_row + 1):
        for column in money_columns:
            ws.cell(
                row,
                column,
            ).number_format = '#,##0.00'

    finalize_sheet(ws)

    return ws


# ============================================================
# EXCEL - CUSTOMERS
# ============================================================

def build_customers_sheet(
    workbook,
    customers=None,
    title="Khách hàng",
):
    """
    Create customer worksheet.
    """

    ws = workbook.create_sheet(title)

    headers = [
        "Mã khách hàng",
        "Tên khách hàng",
        "Số điện thoại",
        "Ngày sinh",
        "Màu yêu thích",
        "Hoa yêu thích",
        "Dịp",
        "Ghi chú",
        "Số đơn hàng",
        "Tổng tiền mua",
        "Lần mua gần nhất",
    ]

    prepare_sheet(
        ws,
        headers,
    )

    if customers is None:
        customers = (
            Customer.query
            .order_by(Customer.id.asc())
            .all()
        )

    for customer in customers:

        orders = customer.orders or []

        order_count = len(orders)

        total_purchase = sum(
            (
                Decimal(
                    str(order.total or 0)
                )
                for order in orders
            ),
            Decimal("0"),
        )

        latest_order_date = None

        for order in orders:

            if not order.created:
                continue

            if (
                latest_order_date is None
                or order.created
                > latest_order_date
            ):
                latest_order_date = (
                    order.created
                )

        ws.append([
            customer.id,
            customer.name,
            customer.phone,
            format_date(
                customer.birthday
            ),
            customer.colors or "",
            customer.flowers or "",
            customer.occasions or "",
            customer.notes or "",
            order_count,
            float(total_purchase),
            format_date(
                latest_order_date
            ),
        ])

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row,
            4,
        ).number_format = "dd/mm/yyyy"

        ws.cell(
            row,
            11,
        ).number_format = "dd/mm/yyyy"

        ws.cell(
            row,
            10,
        ).number_format = '#,##0.00'

    finalize_sheet(ws)

    return ws


# ============================================================
# EXCEL - INVENTORY
# ============================================================

def build_inventory_sheet(
    workbook,
    inventory=None,
    title="Tồn kho",
):
    """
    Create inventory worksheet.
    """

    ws = workbook.create_sheet(title)

    headers = [
        "Mã vật tư",
        "Tên vật tư",
        "Danh mục",
        "Đơn vị",
        "Tồn kho",
        "Tồn tối thiểu",
        "Đơn giá vốn",
        "Giá trị tồn kho",
        "Trạng thái tồn kho",
    ]

    prepare_sheet(
        ws,
        headers,
    )

    if inventory is None:
        inventory = (
            Inventory.query
            .order_by(Inventory.id.asc())
            .all()
        )

    for item in inventory:

        stock = Decimal(
            str(item.stock or 0)
        )

        min_stock = Decimal(
            str(item.min_stock or 0)
        )

        cost = Decimal(
            str(item.cost or 0)
        )

        stock_value = (
            stock * cost
        )

        if stock <= 0:
            stock_status = "Hết hàng"

        elif stock <= min_stock:
            stock_status = "Sắp hết"

        else:
            stock_status = "Bình thường"

        ws.append([
            item.id,
            item.name,
            item.category or "",
            item.unit or "",
            float(stock),
            float(min_stock),
            float(cost),
            float(stock_value),
            stock_status,
        ])

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row,
            5,
        ).number_format = '#,##0.000'

        ws.cell(
            row,
            6,
        ).number_format = '#,##0.000'

        ws.cell(
            row,
            7,
        ).number_format = '#,##0.00'

        ws.cell(
            row,
            8,
        ).number_format = '#,##0.00'

    finalize_sheet(ws)

    return ws


# ============================================================
# EXCEL - MOVEMENTS
# ============================================================

def build_movements_sheet(
    workbook,
    movements=None,
    title="Biến động kho",
):
    """
    Create inventory movement worksheet.
    """

    ws = workbook.create_sheet(title)

    headers = [
        "Mã biến động",
        "Ngày",
        "Mã vật tư",
        "Tên vật tư",
        "Loại",
        "Số lượng",
        "Đơn giá",
        "Thành tiền",
        "Ghi chú",
    ]

    prepare_sheet(
        ws,
        headers,
    )

    if movements is None:
        movements = (
            Movement.query
            .order_by(
                Movement.movement_date.desc(),
                Movement.id.desc(),
            )
            .all()
        )

    for movement in movements:

        qty = Decimal(
            str(movement.qty or 0)
        )

        unit_cost = Decimal(
            str(movement.unit_cost or 0)
        )

        total = (
            qty * unit_cost
        )

        inventory_item = (
            movement.inventory_item
        )

        ws.append([
            movement.id,

            format_date(
                movement.movement_date
            ),

            movement.item_id,

            inventory_item.name
            if inventory_item
            else "",

            MOVEMENT_TYPES.get(
                movement.type,
                movement.type or "",
            ),

            float(qty),
            float(unit_cost),
            float(total),

            movement.note or "",
        ])

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row,
            2,
        ).number_format = "dd/mm/yyyy"

        ws.cell(
            row,
            6,
        ).number_format = '#,##0.000'

        ws.cell(
            row,
            7,
        ).number_format = '#,##0.00'

        ws.cell(
            row,
            8,
        ).number_format = '#,##0.00'

    finalize_sheet(ws)

    return ws


# ============================================================
# EXCEL - WORKFLOW
# ============================================================

def build_workflow_sheet(
    workbook,
    orders=None,
    title="Sản xuất & giao",
):
    """
    Create production / delivery worksheet.
    """

    ws = workbook.create_sheet(title)

    headers = [
        "Mã đơn hàng",
        "Ngày tạo",

        "Mã khách hàng",
        "Tên khách hàng",
        "SĐT khách hàng",

        "Người nhận",
        "SĐT người nhận",

        "Sản phẩm",

        "Ngày giao",
        "Giờ giao",

        "Địa chỉ",
        "Trạng thái",

        "Tổng tiền",
        "Đã cọc",
        "Còn phải thu",

        "Ghi chú",
    ]

    prepare_sheet(
        ws,
        headers,
    )

    if orders is None:
        orders = (
            Order.query
            .order_by(
                Order.created.desc(),
                Order.id.desc(),
            )
            .all()
        )

    for order in orders:

        customer = order.customer

        ws.append([
            order.id,
            format_date(
                order.created
            ),

            customer.id
            if customer
            else order.customer_id,

            customer.name
            if customer
            else "",

            customer.phone
            if customer
            else "",

            order.recipient or "",
            order.recipient_phone or "",

            order.product or "",

            format_date(
                order.delivery_date
            ),

            order.delivery_time or "",

            order.address or "",

            order.status or "",

            format_number(
                order.total
            ),

            format_number(
                order.deposit
            ),

            format_number(
                order.balance
            ),

            order.notes or "",
        ])

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row,
            2,
        ).number_format = "dd/mm/yyyy"

        ws.cell(
            row,
            9,
        ).number_format = "dd/mm/yyyy"

        for column in [
            13,
            14,
            15,
        ]:
            ws.cell(
                row,
                column,
            ).number_format = '#,##0.00'

    finalize_sheet(ws)

    return ws


# ============================================================
# EXCEL - SUMMARY
# ============================================================

def build_summary_sheet(
    workbook,
    customers,
    inventory,
    movements,
    orders,
):
    """
    Create overview sheet.
    """

    ws = workbook.create_sheet(
        "Tong_Quan",
        0,
    )

    ws["A1"] = "LINCHAU FLORAL - TỔNG QUAN DỮ LIỆU"

    ws["A1"].font = Font(
        bold=True,
        size=16,
    )

    ws["A3"] = "Chỉ tiêu"
    ws["B3"] = "Giá trị"

    style_header(
        ws,
        row=3,
    )

    total_order_value = sum(
        (
            Decimal(
                str(order.total or 0)
            )
            for order in orders
        ),
        Decimal("0"),
    )

    total_deposit = sum(
        (
            Decimal(
                str(order.deposit or 0)
            )
            for order in orders
        ),
        Decimal("0"),
    )

    total_balance = sum(
        (
            Decimal(
                str(order.balance or 0)
            )
            for order in orders
        ),
        Decimal("0"),
    )

    total_inventory_value = sum(
        (
            Decimal(
                str(item.stock or 0)
            )
            * Decimal(
                str(item.cost or 0)
            )
            for item in inventory
        ),
        Decimal("0"),
    )

    rows = [
        (
            "Số khách hàng",
            len(customers),
        ),

        (
            "Số đơn hàng",
            len(orders),
        ),

        (
            "Số mặt hàng tồn kho",
            len(inventory),
        ),

        (
            "Số biến động kho",
            len(movements),
        ),

        (
            "Tổng giá trị đơn hàng",
            float(total_order_value),
        ),

        (
            "Tổng tiền đã cọc",
            float(total_deposit),
        ),

        (
            "Tổng còn phải thu",
            float(total_balance),
        ),

        (
            "Tổng giá trị tồn kho",
            float(total_inventory_value),
        ),
    ]

    for row in rows:
        ws.append(row)

    for row in range(
        4,
        ws.max_row + 1,
    ):
        ws.cell(
            row,
            2,
        ).alignment = Alignment(
            horizontal="right"
        )

    for row in range(8, 12):
        ws.cell(
            row,
            2,
        ).number_format = '#,##0.00'

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 22

    return ws


# ============================================================
# API - EXPORT ORDERS
# ============================================================

@export_bp.get("/api/export/orders")
def export_orders():
    """
    Export orders to Excel.

    Optional filters:

        ?q=KH001
        ?status=Đã giao
        ?date_from=2026-09-01
        ?date_to=2026-09-30
    """

    try:

        orders = get_filtered_orders()

        workbook = Workbook()

        # Remove default sheet
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        build_orders_sheet(
            workbook,
            orders,
        )

        return workbook_to_file(
            workbook,
            "Don_Hang.xlsx",
        )

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 500


# ============================================================
# API - EXPORT CUSTOMERS
# ============================================================
@export_bp.get("/api/export/customers")
def export_customers():

    try:

        query = (
            Customer.query
        )

        # ====================================================
        # Keyword
        # ====================================================

        keyword = (
            request.args.get("q", "")
            .strip()
            .lower()
        )

        if keyword:
            pattern = f"%{keyword}%"

            query = query.filter(
                or_(
                    func.lower(
                        Customer.id
                    ).like(pattern),

                    func.lower(
                        Customer.name
                    ).like(pattern),

                    func.lower(
                        Customer.phone
                    ).like(pattern),

                    func.lower(
                        Customer.colors
                    ).like(pattern),

                    func.lower(
                        Customer.flowers
                    ).like(pattern),

                    func.lower(
                        Customer.occasions
                    ).like(pattern),

                    func.lower(
                        Customer.notes
                    ).like(pattern),
                )
            )

        customers = (
            query
            .order_by(Customer.id.asc())
            .all()
        )

        workbook = Workbook()

        default_sheet = workbook.active
        workbook.remove(default_sheet)

        build_customers_sheet(
            workbook,
            customers,
        )

        filename = (
            "Khach_Hang_Dang_Loc.xlsx"
            if keyword
            else "Khach_Hang.xlsx"
        )

        return workbook_to_file(
            workbook,
            filename,
        )

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 500

# ============================================================
# API - EXPORT INVENTORY
# ============================================================
@export_bp.get("/api/export/inventory")
def export_inventory():

    try:

        query = (
            Inventory.query
        )

        # ====================================================
        # Keyword
        # ====================================================

        keyword = (
            request.args.get("q", "")
            .strip()
            .lower()
        )

        if keyword:
            pattern = f"%{keyword}%"

            query = query.filter(
                or_(
                    func.lower(
                        Inventory.id
                    ).like(pattern),

                    func.lower(
                        Inventory.name
                    ).like(pattern),

                    func.lower(
                        Inventory.category
                    ).like(pattern),

                    func.lower(
                        Inventory.unit
                    ).like(pattern),
                )
            )

        inventory = (
            query
            .order_by(Inventory.id.asc())
            .all()
        )

        workbook = Workbook()

        default_sheet = workbook.active
        workbook.remove(default_sheet)

        build_inventory_sheet(
            workbook,
            inventory,
        )

        filename = (
            "Kho_Dang_Loc.xlsx"
            if keyword
            else "Kho.xlsx"
        )

        return workbook_to_file(
            workbook,
            filename,
        )

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 500
    
# ============================================================
# API - EXPORT MOVEMENTS
# ============================================================
@export_bp.get("/api/export/movements")
def export_movements():

    try:

        query = (
            Movement.query
            .outerjoin(
                Inventory,
                Movement.item_id == Inventory.id,
            )
        )

        # ====================================================
        # Keyword
        # ====================================================

        keyword = (
            request.args.get("q", "")
            .strip()
            .lower()
        )

        if keyword:
            pattern = f"%{keyword}%"

            query = query.filter(
                or_(
                    func.lower(
                        Movement.id
                    ).like(pattern),

                    func.lower(
                        Movement.item_id
                    ).like(pattern),

                    func.lower(
                        Inventory.name
                    ).like(pattern),

                    func.lower(
                        Movement.note
                    ).like(pattern),
                )
            )

        # ====================================================
        # Movement type
        # ====================================================

        movement_type = (
            request.args.get("type", "")
            .strip()
            .lower()
        )

        if movement_type:
            query = query.filter(
                Movement.type == movement_type
            )

        # ====================================================
        # Order
        # ====================================================

        movements = (
            query
            .order_by(
                Movement.movement_date.desc(),
                Movement.id.desc(),
            )
            .all()
        )

        workbook = Workbook()

        default_sheet = workbook.active
        workbook.remove(default_sheet)

        build_movements_sheet(
            workbook,
            movements,
        )

        filename = (
            "Bien_Dong_Kho_Dang_Loc.xlsx"
            if keyword or movement_type
            else "Bien_Dong_Kho.xlsx"
        )

        return workbook_to_file(
            workbook,
            filename,
        )

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 500
    
# ============================================================
# API - EXPORT WORKFLOW
# ============================================================

@export_bp.get("/api/export/workflow")
def export_workflow():

    try:

        orders = (
            Order.query
            .order_by(
                Order.created.desc(),
                Order.id.desc(),
            )
            .all()
        )

        workbook = Workbook()

        default_sheet = workbook.active
        workbook.remove(default_sheet)

        build_workflow_sheet(
            workbook,
            orders,
        )

        return workbook_to_file(
            workbook,
            "San_Xuat_Giao.xlsx",
        )

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 500


# ============================================================
# API - EXPORT ALL
# ============================================================

@export_bp.get("/api/export/all")
def export_all():

    try:

        customers = (
            Customer.query
            .order_by(Customer.id.asc())
            .all()
        )

        inventory = (
            Inventory.query
            .order_by(Inventory.id.asc())
            .all()
        )

        movements = (
            Movement.query
            .order_by(
                Movement.movement_date.desc(),
                Movement.id.desc(),
            )
            .all()
        )

        orders = (
            Order.query
            .order_by(
                Order.created.desc(),
                Order.id.desc(),
            )
            .all()
        )

        workbook = Workbook()

        default_sheet = workbook.active
        workbook.remove(default_sheet)

        # ----------------------------------------------------
        # Summary
        # ----------------------------------------------------

        build_summary_sheet(
            workbook,
            customers,
            inventory,
            movements,
            orders,
        )

        # ----------------------------------------------------
        # Orders
        # ----------------------------------------------------

        build_orders_sheet(
            workbook,
            orders,
            title="Don_Hang",
        )

        # ----------------------------------------------------
        # Customers
        # ----------------------------------------------------

        build_customers_sheet(
            workbook,
            customers,
            title="Khach_Hang",
        )

        # ----------------------------------------------------
        # Inventory
        # ----------------------------------------------------

        build_inventory_sheet(
            workbook,
            inventory,
            title="Ton_Kho",
        )

        # ----------------------------------------------------
        # Movements
        # ----------------------------------------------------

        build_movements_sheet(
            workbook,
            movements,
            title="Bien_Dong_Kho",
        )

        # ----------------------------------------------------
        # Workflow
        # ----------------------------------------------------

        build_workflow_sheet(
            workbook,
            orders,
            title="San_Xuat_Giao",
        )

        return workbook_to_file(
            workbook,
            "LinChau_Full_Data.xlsx",
        )

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 500


# ============================================================
# BACKUP - SERIALIZE CUSTOMER
# ============================================================

def serialize_customer(customer):

    return OrderedDict([
        (
            "id",
            customer.id,
        ),

        (
            "name",
            customer.name,
        ),

        (
            "phone",
            customer.phone,
        ),

        (
            "birthday",
            safe_value(
                customer.birthday
            ),
        ),

        (
            "colors",
            customer.colors,
        ),

        (
            "flowers",
            customer.flowers,
        ),

        (
            "occasions",
            customer.occasions,
        ),

        (
            "notes",
            customer.notes,
        ),
    ])


# ============================================================
# BACKUP - SERIALIZE INVENTORY
# ============================================================

def serialize_inventory(item):

    return OrderedDict([
        (
            "id",
            item.id,
        ),

        (
            "name",
            item.name,
        ),

        (
            "category",
            item.category,
        ),

        (
            "unit",
            item.unit,
        ),

        (
            "stock",
            safe_value(item.stock),
        ),

        (
            "min_stock",
            safe_value(item.min_stock),
        ),

        (
            "cost",
            safe_value(item.cost),
        ),
    ])


# ============================================================
# BACKUP - SERIALIZE MOVEMENT
# ============================================================

def serialize_movement(movement):

    return OrderedDict([
        (
            "id",
            movement.id,
        ),

        (
            "movement_date",
            safe_value(
                movement.movement_date
            ),
        ),

        (
            "item_id",
            movement.item_id,
        ),

        (
            "type",
            movement.type,
        ),

        (
            "qty",
            safe_value(
                movement.qty
            ),
        ),

        (
            "unit_cost",
            safe_value(
                movement.unit_cost
            ),
        ),

        (
            "note",
            movement.note,
        ),
    ])


# ============================================================
# BACKUP - SERIALIZE ORDER
# ============================================================

def serialize_order(order):

    return OrderedDict([
        (
            "id",
            order.id,
        ),

        (
            "created",
            safe_value(
                order.created
            ),
        ),

        (
            "customer_id",
            order.customer_id,
        ),

        (
            "recipient",
            order.recipient,
        ),

        (
            "recipient_phone",
            order.recipient_phone,
        ),

        (
            "delivery_date",
            safe_value(
                order.delivery_date
            ),
        ),

        (
            "delivery_time",
            order.delivery_time,
        ),

        (
            "address",
            order.address,
        ),

        (
            "product",
            order.product,
        ),

        (
            "subtotal",
            safe_value(
                order.subtotal
            ),
        ),

        (
            "shipping",
            safe_value(
                order.shipping
            ),
        ),

        (
            "discount_type",
            order.discount_type,
        ),

        (
            "discount_percent",
            safe_value(
                order.discount_percent
            ),
        ),

        (
            "discount_amount",
            safe_value(
                order.discount_amount
            ),
        ),

        (
            "discount_value",
            safe_value(
                order.discount_value
            ),
        ),

        (
            "total",
            safe_value(
                order.total
            ),
        ),

        (
            "deposit",
            safe_value(
                order.deposit
            ),
        ),

        (
            "balance",
            safe_value(
                order.balance
            ),
        ),

        (
            "payment",
            order.payment,
        ),

        (
            "status",
            order.status,
        ),

        (
            "card",
            order.card,
        ),

        (
            "notes",
            order.notes,
        ),
    ])


# ============================================================
# API - BACKUP JSON
# ============================================================

@export_bp.get("/api/backup")
def backup_data():

    try:

        customers = (
            Customer.query
            .order_by(Customer.id.asc())
            .all()
        )

        inventory = (
            Inventory.query
            .order_by(Inventory.id.asc())
            .all()
        )

        movements = (
            Movement.query
            .order_by(Movement.id.asc())
            .all()
        )

        orders = (
            Order.query
            .order_by(Order.id.asc())
            .all()
        )

        backup = OrderedDict([
            (
                "app",
                "LinChau Floral",
            ),

            (
                "backup_version",
                "1.0",
            ),

            (
                "created_at",
                datetime.now().isoformat(),
            ),

            (
                "tables",
                OrderedDict([
                    (
                        "customers",
                        [
                            serialize_customer(
                                customer
                            )
                            for customer in customers
                        ],
                    ),

                    (
                        "inventory",
                        [
                            serialize_inventory(
                                item
                            )
                            for item in inventory
                        ],
                    ),

                    (
                        "movements",
                        [
                            serialize_movement(
                                movement
                            )
                            for movement in movements
                        ],
                    ),

                    (
                        "orders",
                        [
                            serialize_order(
                                order
                            )
                            for order in orders
                        ],
                    ),
                ]),
            ),
        ])

        filename = (
            "LinChau_Backup_"
            f"{datetime.now().strftime('%Y%m%d_%H%M%S')}"
            ".json"
        )

        return json_file_response(
            backup,
            filename,
        )

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 500


# ============================================================
# RESTORE HELPERS
# ============================================================

def parse_date(value):
    """
    Convert YYYY-MM-DD string to date.
    """
    if value in (
        None,
        "",
    ):
        return None

    if isinstance(value, date):
        return value

    return datetime.strptime(
        str(value),
        "%Y-%m-%d",
    ).date()


def parse_decimal(value):
    """
    Convert JSON numeric value to Decimal.
    """
    if value in (
        None,
        "",
    ):
        return Decimal("0")

    return Decimal(
        str(value)
    )


def restore_customer(data):

    return Customer(
        id=str(
            data.get("id", "")
        ),

        name=data.get(
            "name",
            "",
        ),

        phone=data.get(
            "phone",
            "",
        ),

        birthday=parse_date(
            data.get("birthday")
        ),

        colors=data.get(
            "colors"
        ),

        flowers=data.get(
            "flowers"
        ),

        occasions=data.get(
            "occasions"
        ),

        notes=data.get(
            "notes"
        ),
    )


def restore_inventory(data):

    return Inventory(
        id=str(
            data.get("id", "")
        ),

        name=data.get(
            "name",
            "",
        ),

        category=data.get(
            "category"
        ),

        unit=data.get(
            "unit"
        ),

        stock=parse_decimal(
            data.get("stock")
        ),

        min_stock=parse_decimal(
            data.get("min_stock")
        ),

        cost=parse_decimal(
            data.get("cost")
        ),
    )


def restore_movement(data):

    return Movement(
        id=str(
            data.get("id", "")
        ),

        movement_date=parse_date(
            data.get(
                "movement_date"
            )
        ),

        item_id=str(
            data.get(
                "item_id",
                "",
            )
        ),

        type=data.get(
            "type",
            "",
        ),

        qty=parse_decimal(
            data.get("qty")
        ),

        unit_cost=parse_decimal(
            data.get(
                "unit_cost"
            )
        ),

        note=data.get(
            "note"
        ),
    )


def restore_order(data):

    return Order(
        id=str(
            data.get("id", "")
        ),

        created=parse_date(
            data.get("created")
        ),

        customer_id=str(
            data.get(
                "customer_id",
                "",
            )
        ),

        recipient=data.get(
            "recipient"
        ),

        recipient_phone=data.get(
            "recipient_phone"
        ),

        delivery_date=parse_date(
            data.get(
                "delivery_date"
            )
        ),

        delivery_time=data.get(
            "delivery_time"
        ),

        address=data.get(
            "address"
        ),

        product=data.get(
            "product"
        ),

        subtotal=parse_decimal(
            data.get("subtotal")
        ),

        shipping=parse_decimal(
            data.get("shipping")
        ),

        discount_type=data.get(
            "discount_type"
        ),

        discount_percent=parse_decimal(
            data.get(
                "discount_percent"
            )
        ),

        discount_amount=parse_decimal(
            data.get(
                "discount_amount"
            )
        ),

        discount_value=parse_decimal(
            data.get(
                "discount_value"
            )
        ),

        total=parse_decimal(
            data.get("total")
        ),

        deposit=parse_decimal(
            data.get("deposit")
        ),

        balance=parse_decimal(
            data.get("balance")
        ),

        payment=data.get(
            "payment"
        ),

        status=data.get(
            "status",
            "Chờ xác nhận",
        ),

        card=data.get(
            "card"
        ),

        notes=data.get(
            "notes"
        ),
    )


# ============================================================
# VALIDATE BACKUP
# ============================================================

def validate_backup(data):

    if not isinstance(data, dict):
        raise ValueError(
            "File backup không có định dạng JSON hợp lệ."
        )

    if data.get("app") != "LinChau Floral":
        raise ValueError(
            "File không phải backup của LinChau Floral."
        )

    tables = data.get("tables")

    if not isinstance(tables, dict):
        raise ValueError(
            "Backup không có phần tables."
        )

    required_tables = [
        "customers",
        "inventory",
        "movements",
        "orders",
    ]

    for table_name in required_tables:

        if table_name not in tables:
            raise ValueError(
                f"Backup thiếu bảng: {table_name}"
            )

        if not isinstance(
            tables[table_name],
            list,
        ):
            raise ValueError(
                f"Dữ liệu bảng {table_name} không hợp lệ."
            )

    return True


# ============================================================
# API - RESTORE JSON
# ============================================================

@export_bp.post("/api/restore")
def restore_data():

    try:

        # ----------------------------------------------------
        # Read JSON
        # ----------------------------------------------------

        data = request.get_json(
            silent=True
        )

        if not data:
            return jsonify({
                "success": False,
                "message": (
                    "Không nhận được dữ liệu JSON."
                ),
            }), 400

        validate_backup(data)

        tables = data["tables"]

        customers_data = tables[
            "customers"
        ]

        inventory_data = tables[
            "inventory"
        ]

        movements_data = tables[
            "movements"
        ]

        orders_data = tables[
            "orders"
        ]

        # ----------------------------------------------------
        # Validate foreign keys before changing DB
        # ----------------------------------------------------

        customer_ids = {
            str(item.get("id"))
            for item in customers_data
        }

        inventory_ids = {
            str(item.get("id"))
            for item in inventory_data
        }

        for order_data in orders_data:

            customer_id = str(
                order_data.get(
                    "customer_id",
                    "",
                )
            )

            if customer_id not in customer_ids:
                raise ValueError(
                    "Đơn hàng "
                    f"{order_data.get('id')} "
                    "tham chiếu đến khách hàng "
                    f"{customer_id} không tồn tại "
                    "trong backup."
                )

        for movement_data in movements_data:

            item_id = str(
                movement_data.get(
                    "item_id",
                    "",
                )
            )

            if item_id not in inventory_ids:
                raise ValueError(
                    "Biến động kho "
                    f"{movement_data.get('id')} "
                    "tham chiếu đến vật tư "
                    f"{item_id} không tồn tại "
                    "trong backup."
                )

        # ----------------------------------------------------
        # Transaction
        # ----------------------------------------------------

        try:

            # Delete child tables first
            Order.query.delete(
                synchronize_session=False
            )

            Movement.query.delete(
                synchronize_session=False
            )

            # Then parent tables
            Customer.query.delete(
                synchronize_session=False
            )

            Inventory.query.delete(
                synchronize_session=False
            )

            db.session.flush()

            # ------------------------------------------------
            # Customers
            # ------------------------------------------------

            for item in customers_data:

                customer = restore_customer(
                    item
                )

                db.session.add(
                    customer
                )

            db.session.flush()

            # ------------------------------------------------
            # Inventory
            # ------------------------------------------------

            for item in inventory_data:

                inventory = restore_inventory(
                    item
                )

                db.session.add(
                    inventory
                )

            db.session.flush()

            # ------------------------------------------------
            # Movements
            # ------------------------------------------------

            for item in movements_data:

                movement = restore_movement(
                    item
                )

                db.session.add(
                    movement
                )

            db.session.flush()

            # ------------------------------------------------
            # Orders
            # ------------------------------------------------

            for item in orders_data:

                order = restore_order(
                    item
                )

                db.session.add(
                    order
                )

            db.session.commit()

        except Exception:

            db.session.rollback()

            raise

        return jsonify({
            "success": True,
            "message": (
                "Khôi phục dữ liệu thành công."
            ),
            "counts": {
                "customers": len(
                    customers_data
                ),
                "inventory": len(
                    inventory_data
                ),
                "movements": len(
                    movements_data
                ),
                "orders": len(
                    orders_data
                ),
            },
        })

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 400


# ============================================================
# DATA HEALTH
# ============================================================

@export_bp.get("/api/data-health")
def data_health():

    try:

        customer_count = (
            Customer.query.count()
        )

        inventory_count = (
            Inventory.query.count()
        )

        movement_count = (
            Movement.query.count()
        )

        order_count = (
            Order.query.count()
        )

        # ----------------------------------------------------
        # Orders without valid customer
        # ----------------------------------------------------

        orphan_orders = (
            db.session.query(
                func.count(Order.id)
            )
            .outerjoin(
                Customer,
                Order.customer_id
                == Customer.id,
            )
            .filter(
                Customer.id.is_(None)
            )
            .scalar()
            or 0
        )

        # ----------------------------------------------------
        # Movements without inventory
        # ----------------------------------------------------

        orphan_movements = (
            db.session.query(
                func.count(Movement.id)
            )
            .outerjoin(
                Inventory,
                Movement.item_id
                == Inventory.id,
            )
            .filter(
                Inventory.id.is_(None)
            )
            .scalar()
            or 0
        )

        # ----------------------------------------------------
        # Negative stock
        # ----------------------------------------------------

        negative_stock = (
            Inventory.query
            .filter(
                Inventory.stock < 0
            )
            .count()
        )

        # ----------------------------------------------------
        # Low stock
        # ----------------------------------------------------

        low_stock = (
            Inventory.query
            .filter(
                Inventory.stock
                <= Inventory.min_stock
            )
            .count()
        )

        # ----------------------------------------------------
        # Orders without delivery date
        # ----------------------------------------------------

        orders_without_delivery = (
            Order.query
            .filter(
                Order.delivery_date.is_(None)
            )
            .count()
        )

        # ----------------------------------------------------
        # Result
        # ----------------------------------------------------

        healthy = (
            orphan_orders == 0
            and orphan_movements == 0
        )

        return jsonify({
            "success": True,

            "healthy": healthy,

            "counts": {
                "customers": customer_count,
                "inventory": inventory_count,
                "movements": movement_count,
                "orders": order_count,
            },

            "warnings": {
                "orphan_orders": orphan_orders,
                "orphan_movements": orphan_movements,
                "negative_stock": negative_stock,
                "low_stock": low_stock,
                "orders_without_delivery": (
                    orders_without_delivery
                ),
            },
        })

    except Exception as exc:

        return jsonify({
            "success": False,
            "message": str(exc),
        }), 500